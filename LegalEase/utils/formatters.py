import html, re
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

REPLACEMENTS = {"“": '"', "”": '"', "‘": "'", "’": "'", "–": "-", "—": "-", "…": "...", "\u00a0": " "}

def sanitize_text(text: str) -> str:
    for old, new in REPLACEMENTS.items():
        text = text.replace(old, new)
    return text.replace("\x00", "").strip()

def _clean_markdown(line: str) -> str:
    return re.sub(r"[*_`#]", "", line).strip()

def format_html_preview(text: str) -> str:
    safe = html.escape(text)
    safe = re.sub(r"^### (.+)$", r"<h4>\1</h4>", safe, flags=re.M)
    safe = re.sub(r"^## (.+)$", r"<h3>\1</h3>", safe, flags=re.M)
    safe = re.sub(r"^# (.+)$", r"<h2>\1</h2>", safe, flags=re.M)
    safe = safe.replace("\n", "<br>")
    return f'<div style="background:#111827;color:#e5e7eb;padding:22px;border-radius:12px;line-height:1.6;max-height:560px;overflow:auto;border:1px solid #374151">{safe}</div>'

def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7); section.bottom_margin = Inches(0.7)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(doc_type.upper()); r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(16)
    for line in sanitize_text(text).splitlines():
        line = line.strip()
        if not line: continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(_clean_markdown(line)); run.font.name = "Times New Roman"; run.font.size = Pt(11)
        if line.startswith("#") or re.match(r"^\d+[.)]\s", line): run.bold = True
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated with LegalEase - Review with a qualified legal professional before use.").font.size = Pt(8)
    out = BytesIO(); doc.save(out); return out.getvalue()

def format_pdf(text: str, doc_type: str) -> bytes:
    class LegalPDF(FPDF):
        def footer(self):
            self.set_y(-12); self.set_font("Helvetica", size=8)
            self.cell(0, 8, f"LegalEase | Page {self.page_no()} | Review with a qualified legal professional", align="C")
    pdf = LegalPDF(); pdf.set_auto_page_break(True, 18); pdf.add_page()
    pdf.set_font("Helvetica", "B", 16); pdf.multi_cell(0, 9, sanitize_text(doc_type).upper(), align="C"); pdf.ln(3)
    for line in sanitize_text(text).splitlines():
        line = _clean_markdown(line)
        if not line: pdf.ln(2); continue
        heading = bool(re.match(r"^\d+[.)]\s", line))
        pdf.set_font("Helvetica", "B" if heading else "", 11)
        pdf.multi_cell(0, 6, line)
    return bytes(pdf.output())
