import re
import html
from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer


def sanitize_text(text):
    """
    Clean text before displaying or exporting.
    """
    if text is None:
        return ""

    text = str(text)

    # Remove null characters
    text = text.replace("\x00", "")

    # Normalize excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def format_docx(text, doc_type=None):
    doc = Document()

    # Optional title based on selected document type
    if doc_type:
        doc.add_heading(str(doc_type), level=0)

    # Add generated/edited content
    for paragraph in text.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()

def format_pdf(text, doc_type=None):
    from io import BytesIO
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import inch

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "LegalTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=16,
        leading=20,
        spaceAfter=20
    )

    body_style = ParagraphStyle(
        "LegalBody",
        parent=styles["BodyText"],
        fontSize=11,
        leading=16,
        spaceAfter=10
    )

    story = []

    # Optional document title
    if doc_type:
        title = str(doc_type).replace("_", " ").title()
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 0.2 * inch))

    # Convert text into paragraphs
    paragraphs = str(text).split("\n")

    for paragraph in paragraphs:
        paragraph = paragraph.strip()

        if paragraph:
            # Escape characters that may interfere with ReportLab
            paragraph = (
                paragraph
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            story.append(Paragraph(paragraph, body_style))
        else:
            story.append(Spacer(1, 8))

    doc.build(story)

    buffer.seek(0)

    return buffer.getvalue()

def format_html_preview(text):
    """
    Convert plain text into safe HTML for preview.
    """
    text = sanitize_text(text)

    escaped_text = html.escape(text)

    paragraphs = escaped_text.split("\n")

    html_content = ""

    for paragraph in paragraphs:
        if paragraph.strip():
            html_content += f"<p>{paragraph}</p>"
        else:
            html_content += "<br>"

    return html_content